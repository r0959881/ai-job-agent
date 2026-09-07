import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT_COLUMNS = [
    ("Job ID", "id"),
    ("Title", "title"),
    ("Company", "company"),
    ("Location", "location"),
    ("Source", "source"),
    ("Posted At", "posted_at"),
    ("Discovered At", "discovered_at"),
    ("Category", "category"),
    ("Contract Type", "contract_type"),
    ("Salary Min", "salary_min"),
    ("Salary Max", "salary_max"),
    ("Currency", "salary_currency"),
    ("Status", "status"),
    ("Job URL", "url"),
    ("Tailored CV", "cv_path"),
    ("Description", "description"),
]


def find_existing_cv(job, output_directory="output"):
    """Find a previously generated CV when older database rows have no cv_path."""
    saved_path = job.get("cv_path")
    if saved_path and Path(saved_path).exists():
        return saved_path

    company = "".join(c for c in str(job.get("company", "")) if c.isalnum() or c in (" ", "_"))
    title = "".join(c for c in str(job.get("title", "")) if c.isalnum() or c in (" ", "_"))
    prefix = f"Ammar_Haider_CV_{company}_{title}".strip()
    candidates = sorted(Path(output_directory).glob(f"{prefix}*.docx"))
    return str(candidates[0]) if candidates else ""


def export_jobs_to_excel(jobs, output_path, excluded_jobs=None):
    """Write job details and links to the tailored CV files into an Excel workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Jobs"

    headers = [header for header, _ in REPORT_COLUMNS]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for job in jobs:
        job["cv_path"] = find_existing_cv(job)
        values = []
        for header, key in REPORT_COLUMNS:
            value = job.get(key, "") or ""
            if key in {"url", "cv_path"} and value:
                value = str(value)
            values.append(value)
        worksheet.append(values)

        row_number = worksheet.max_row
        url_cell = worksheet.cell(row_number, 14)
        cv_cell = worksheet.cell(row_number, 15)
        if url_cell.value:
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"
        cv_path = Path(str(cv_cell.value)) if cv_cell.value else None
        if cv_path and cv_path.exists():
            # A local URI avoids SharePoint treating the sibling filename as a web route.
            cv_cell.hyperlink = cv_path.resolve().as_uri()
            cv_cell.value = "Open tailored CV"
            cv_cell.style = "Hyperlink"
        else:
            cv_cell.value = "CV unavailable"
            cv_cell.font = Font(color="9C0006", italic=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    widths = {
        1: 34, 2: 35, 3: 28, 4: 28, 5: 22, 6: 22, 7: 28,
        8: 20, 9: 18, 10: 14, 11: 14, 12: 12, 13: 14, 14: 50,
        15: 20, 16: 80,
    }
    for column_number, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_number)].width = width

    summary = workbook.create_sheet("Run Info")
    summary.append(["Generated At (UTC)", datetime.now(timezone.utc).isoformat()])
    summary.append(["Jobs Included", len(jobs)])
    summary.append(["CV links", "Open the linked CV from the Jobs sheet"])
    summary.append(["Excluded jobs", "See the Excluded Jobs sheet for filtered jobs"])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 55

    if excluded_jobs:
        excluded = workbook.create_sheet("Excluded Jobs")
        excluded.append(["Job ID", "Title", "Company", "Location", "Status", "Posted At", "Discovered At", "Reason"])
        for cell in excluded[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for job in excluded_jobs:
            excluded.append([
                job.get("id", ""), job.get("title", ""), job.get("company", ""),
                job.get("location", ""), job.get("status", ""), job.get("posted_at", ""),
                job.get("discovered_at", ""), job.get("exclusion_reason", "Filtered out"),
            ])
        excluded.freeze_panes = "A2"
        excluded.auto_filter.ref = excluded.dimensions

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)
    return output_path
