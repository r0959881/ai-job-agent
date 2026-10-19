"""Import a PDF or DOCX resume: extract text locally, structure it with Groq, return a reviewable draft."""
import io

from api.importers.merge import build_import_result

MIN_TEXT = 200


class ResumeImportError(Exception):
    pass


def extract_text(filename, data):
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeImportError("The PDF is password-protected. Save an unprotected copy and upload that.")
            return "\n".join(page.extract_text() or "" for page in reader.pages[:10]), "pdf"
        except PdfReadError as exc:
            raise ResumeImportError("The PDF could not be read. Try exporting it again or upload a DOCX.") from exc
    if name.endswith(".docx") or data[:2] == b"PK":
        from docx import Document

        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise ResumeImportError("The Word file could not be opened. Save it as .docx and try again.") from exc
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines), "docx"
    raise ResumeImportError("Upload a PDF or a Word (.docx) file.")


def docx_preview(data):
    import mammoth

    from api.services.preview import sanitise

    return sanitise(mammoth.convert_to_html(io.BytesIO(data)).value)


def import_resume_file(filename, data, use_ai=True, client=None):
    from api.importers.ai import extract_resume
    from tailor import GroqConfigError, describe_error

    text, kind = extract_text(filename, data)
    if len(text.strip()) < MIN_TEXT:
        raise ResumeImportError(
            "Almost no text could be read from this file. If it is a scanned PDF, export it as text-based PDF "
            "or DOCX, or import your LinkedIn data instead."
        )
    try:
        draft, low_confidence = extract_resume(text, client)
    except GroqConfigError as exc:
        raise ResumeImportError(
            "Reading a resume file needs a Groq API key. Add it in Settings → Integrations, "
            "or import your LinkedIn export, which works without one."
        ) from exc
    except Exception as exc:
        raise ResumeImportError(f"Groq could not read the resume: {describe_error(exc)}. Try again.") from exc
    warnings = []
    if low_confidence:
        warnings.append(f"{len(low_confidence)} field(s) were hard to read. They are highlighted for you to check.")
    draft.setdefault("personal_info", {})
    return build_import_result("resume_file", _clean(draft), warnings, low_confidence=low_confidence,
                               preview_html=docx_preview(data) if kind == "docx" else None)


def _clean(draft):
    """Drop empty or incomplete entries the model may return so validation never fails on them."""
    draft["experience"] = [e for e in draft.get("experience") or [] if e.get("company") and e.get("role")]
    draft["projects"] = [p for p in draft.get("projects") or [] if p.get("title")]
    draft["certifications"] = [c for c in draft.get("certifications") or [] if c.get("name")]
    draft["languages"] = [lang for lang in draft.get("languages") or [] if str(lang).strip()]
    if not isinstance(draft.get("education"), dict):
        draft["education"] = {}
    return draft
