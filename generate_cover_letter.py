import re
from datetime import date

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import agent_settings
import user_scope
from agent_log import log
from config import OUTPUT_DIR, PROJECT_ROOT, clean_job_title, document_stem, load_resume_file, long_path, place_of


ROLE_LETTERS = {
    "Data Integration": {
        "opening": "I am writing to express my interest in Data Integration opportunities. My experience combines ETL and ELT development, SQL data preparation, source-system integration, data quality, and reporting delivery.",
        "evidence": "In my project work, I integrated IoT sensor and fleet API data into Delta Lake, unified five source systems into Medallion data layers, and designed a Python, MySQL, and Power BI solution under GDPR requirements. I also transformed a 12-week manual reporting workflow into a single-click Power BI and Power Automate solution.",
        "fit": "I would bring a practical understanding of how reliable ingestion, modelling, validation, and reporting work together. I am comfortable working across technical and business teams, documenting data definitions, and translating requirements into useful data solutions.",
        "closing": "I would welcome the opportunity to discuss how my ETL, SQL, cloud data, and data quality experience could support your integration initiatives.",
    },
    "Data Science": {
        "opening": "I am writing to express my interest in Data Science opportunities. My background combines Python, statistical analysis, machine learning, data preparation, visualization, and practical data engineering experience.",
        "evidence": "My project work includes Random Forest analysis of climate and crop-production data, a housing-price prediction pipeline with an R2 of 0.817 on unseen data, Auckland cycling-demand forecasting with Facebook Prophet, fraud detection, and a computer-vision project. Alongside this, I have built reporting pipelines and Power BI dashboards using Databricks and Microsoft Fabric.",
        "fit": "These projects have strengthened my ability to prepare data carefully, evaluate models, interpret results, and communicate findings clearly. My professional experience also taught me to connect analytical work to stakeholder requirements and decision-ready reporting.",
        "closing": "I would welcome the opportunity to discuss how my analytical foundation, curiosity, and hands-on project experience could contribute to your data science team.",
    },
    "Data Engineer": {
        "opening": "I am writing to express my interest in Data Engineering opportunities. I have practical experience with Azure Databricks, PySpark, Delta Lake, Microsoft Fabric, ETL and ELT pipelines, data quality, and CI/CD.",
        "evidence": "I built a pipeline that ingested more than 50,000 daily IoT sensor and fleet API records into Delta Lake, unified five source systems into Medallion data layers, and implemented automated data quality checks. I also created GitHub Actions pipelines for PySpark testing and infrastructure-as-code scripts that reduced manual setup time by 80 percent.",
        "fit": "I enjoy building data platforms that are reliable, observable, and useful to the people who depend on them. My experience collaborating with stakeholders helps me connect technical implementation with clear business requirements and measurable outcomes.",
        "closing": "I would welcome the opportunity to discuss how my cloud data platform, pipeline, and automation experience could support your engineering team.",
    },
    "Data Analyst": {
        "opening": "I am writing to express my interest in Data Analyst opportunities. My experience combines SQL, Power BI, DAX, KPI development, data modelling, reporting automation, and stakeholder communication.",
        "evidence": "I designed interactive Power BI dashboards for more than 50 business users, wrote SQL queries to clean and shape operational data, and translated stakeholder requirements into KPI definitions and reporting roadmaps. I also modelled organizational data and presented executive KPI dashboards directly to the CEO.",
        "fit": "I focus on making complex information understandable and actionable. From improving a 12-week manual reporting process to documenting metric logic and business use cases, I have learned to combine technical accuracy with clear communication and user adoption.",
        "closing": "I would welcome the opportunity to discuss how my analytical, reporting, and stakeholder-facing experience could contribute to your team.",
    },
    "Functional Analyst (IT)": {
        "opening": "I am writing to express my interest in Functional Analyst opportunities in IT. My experience includes requirements gathering, functional specifications, user stories, stakeholder alignment, reporting solutions, and process improvement.",
        "evidence": "At Estée Lauder Companies, I translated strategic goals into functional specifications, user stories, KPI definitions, and reporting roadmaps. I worked as a liaison between business and technical teams, guided users through iterative reviews, and helped deliver a single-click Power BI and Power Automate solution that replaced a 12-week manual workflow.",
        "fit": "I am comfortable asking clarifying questions, documenting processes and data definitions, and keeping delivery connected to the needs of users. My technical foundation in SQL, Python, Power BI, and data platforms also helps me communicate effectively with engineering teams.",
        "closing": "I would welcome the opportunity to discuss how my functional analysis, data, and stakeholder-management experience could support your IT initiatives.",
    },
}


GENERAL_LETTER = {
    "opening": "I am writing to express my interest in this opportunity. My background combines hands-on data work, clear communication, and a practical focus on useful outcomes.",
    "evidence": "My experience and projects are summarised in the attached CV, and I would be glad to walk you through the work that is most relevant to this role.",
    "fit": "I enjoy learning quickly, collaborating with both technical and business colleagues, and turning requirements into reliable results.",
    "closing": "I would welcome the opportunity to discuss how I could contribute to your team.",
}


def role_letters():
    """Editable letters from the settings, falling back to the built-in text.

    The built-in letters describe this project's original owner, so web users only
    get their own letters plus the general one.
    """
    configured = agent_settings.get("cover_letters") or {}
    builtin = {} if user_scope.current() else ROLE_LETTERS
    return {**builtin, **configured}


SAMPLE_LETTER = """I am applying for the Data Engineer role on Axon. The idea of bringing scattered supply chain data into one clear view, so that teams can make better decisions, is exactly the kind of work I want to do. I finished my Bachelor in Data Science in Mechelen, so your office is also close to where I studied, and I would be glad to join a growing product team there.

Over the past year I have built pipelines on Azure Databricks and Microsoft Fabric. At Estée Lauder I designed and supported ETL/ELT pipelines in PySpark and SQL with Bronze-Silver-Gold layers. I tuned the workloads with partitioning and ZORDER, which cut query latency by 40% and storage costs by 25%, and I set up retries and SLA monitoring in Databricks Workflows.

Your bonus skill, supply chain data, is where I have been practising. I built a pipeline that reconciles ERP orders, MES production batches and WMS warehouse movements into an on-time-in-full KPI, and I found the real problems that appear when systems disagree: duplicates, missing values and broken links between sources.

I want to be open about where I am. I have about six months of professional experience plus these projects, not two to three years. My CI/CD experience is on the CI side so far, with GitHub Actions, and I have not built graph data models or worked with Scala. I learn fast, I like understanding the business side of the data, and I would welcome the mentorship you describe.

Your values, dig deep, own it, come together, move fast and be kind, fit how I like to work. I take ownership, I ask for feedback and I enjoy working with analysts, engineers and the business. Thank you for your time. I would be glad to talk about how I can contribute to Axon."""

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x65, 0x73)
FONT = "Calibri"



def fallback_paragraphs(role_family, data):
    """The letter without Groq: the role family's written letter (or the general one)."""
    letter = role_letters().get(role_family, GENERAL_LETTER)
    role_profile = (data.get("role_profiles") or {}).get(role_family, {})
    return [letter["opening"], letter["evidence"],
            f"My focus is on {role_profile.get('focus', 'turning reliable data into useful outcomes')}. {letter['fit']}",
            letter["closing"]]


def resume_facts(data):
    """The resume as one fact per line, each labelled with the role or project it belongs to.

    With the source on every line the model keeps each result with its own role or project
    (a plain YAML dump let "8 minutes to 45 seconds" drift from a project to a job).
    """
    lines = []
    for entry in data.get("experience") or []:
        if isinstance(entry, dict):
            label = f"[Job: {entry.get('role', '')} at {entry.get('company', '')}, {entry.get('dates', '')}]"
            lines += [f"{label} {bullet}" for bullet in entry.get("bullets") or []]
    for project in data.get("projects") or []:
        if isinstance(project, dict):
            label = f"[Project: {project.get('title', '')} ({project.get('tech_stack', '')})]"
            lines += [f"{label} {highlight}" for highlight in project.get("highlights") or []]
    education = data.get("education") or []
    for entry in [education] if isinstance(education, dict) else education:
        if isinstance(entry, dict):
            lines.append(f"[Education] {entry.get('degree', '')}, {entry.get('institution', '')}, "
                         f"{entry.get('location', '')}, {entry.get('dates', '')}")
    for cert in data.get("certifications") or []:
        lines.append(f"[Certification] {cert.get('name', '') if isinstance(cert, dict) else cert}")
    skills = [str(s) for group in (data.get("core_skills") or {}).values() for s in (group or [])]
    if skills:
        lines.append(f"[Skills] {', '.join(skills)}")
    if data.get("languages"):
        lines.append(f"[Languages] {'; '.join(str(lang) for lang in data['languages'])}")
    return "\n".join(f"    - {line}" for line in lines)


def written_paragraphs(job, data, client=None):
    """A letter for this job, written by Groq in the sample letter's structure. Raises on API errors."""
    from tailor import MODEL, drop_unverifiable_sentences, get_client

    company = job.get("company") or ""
    # Alert jobs carry the alert line (title, extras, tracking link) as description and location.
    description = re.sub(r"https?://\S+", "", job.get("description") or "").strip()
    title = clean_job_title(job.get("title"), company)
    place = place_of(job.get("location"), job.get("title"))
    # Job-alert emails only carry the title line, so there are no requirements to match or to fall short of.
    short_posting = len(description) < 1200
    structure = """
    1. Why this role and this company, using only what the posting says, in plain words.
    2. My most relevant hands-on work: two or three results from my resume, told as one story.
    3. How I work with people, shown by facts from my resume (working with stakeholders, documenting, guiding
       users, presenting), and what that means for this role.
    4. Thanks and a short closing line.""" if short_posting else """
    1. Why this role and this company: what in the posting appeals to me, in plain words.
    2. My most relevant hands-on work, with the concrete results from my resume.
    3. One requirement or bonus skill the posting names, matched with a project or task from my resume.
    4. An honest note on one or two requirements the posting names that my resume does not show. State the gap
       plainly; do not promise courses, plans or timelines.
    5. The company's values or way of working (only if the posting names them), then thanks and a closing line."""
    prompt = f"""
    Write the body of a cover letter for this job, in English (also when the posting is in another language).

    Job title: {title}
    Company: {company or "(unknown: do not name a company)"}
    Location: {place or "(not given)"}
    Job posting:
    \"\"\"
    {description[:6000] or "(no description: rely on the title)"}
    \"\"\"

    My resume, one fact per line, each labelled with the job or project it belongs to (the only source of facts
    about me):
{resume_facts(data)}

    Follow the tone of this example letter, written for a different job at a different company. Take ONLY its
    tone and sentence style from it. Every fact in it (Axon, supply chain, Scala, graph models, values, numbers)
    belongs to that other job: never repeat any of them unless my resume or THIS posting states it.
    \"\"\"
    {SAMPLE_LETTER}
    \"\"\"

    The structure, one paragraph each:{structure}

    RULES:
    - {"4 paragraphs, 230 to 300 words" if short_posting else "4 or 5 paragraphs, 280 to 380 words"} in total.
      Short, plain sentences in the first person, no clichés.
    - Only say the posting asks for, highlights or values something when the posting text above says so. Say
      nothing about the company (its reputation, clients, products or culture) that the posting does not say.
    - Never invent experience, skills, numbers, tools, companies, degrees, plans or courses that are not in my resume.
    - Never mention salary, pay, benefits or contract terms.
    - Use a fact only for the job or project in its label, and never combine facts from different labels into one
      achievement. Describe a project only with the details its facts state. Copy numbers exactly as written.
    - The labels are for you: never copy them. Name the source naturally ("At Estée Lauder, ...",
      "In my Refrigerator Ships project, ...").
    - Pick the two to four strongest facts for this job; do not list everything. At most 90 words per paragraph.
    - Write it as a letter, not a list: vary how sentences start and name each job or project once per paragraph.
    - No greeting, no sign-off, no headings, no bullet points, no markdown: just the paragraphs, separated by a blank line.
    """
    response = (client or get_client()).chat.completions.create(
        model=MODEL, messages=[{"role": "user", "content": prompt}], temperature=0.2)
    text = (response.choices[0].message.content or "").replace("**", "")
    text = re.sub(r"\[(?:Job|Project|Education|Certification|Skills|Languages)\b[^\]]*\]\s*", "", text)
    text = re.sub(r"(\d)[  ]+(?=[%Kk+])", r"\1", text).replace(" ", " ").replace(" ", " ")
    if "\n\n" not in text.strip():  # paragraphs separated by single line breaks
        text = text.replace("\n", "\n\n")
    paragraphs = [drop_unverifiable_sentences(" ".join(part.split()), data) for part in text.split("\n\n")]
    paragraphs = [part for part in paragraphs if part and not part.lower().startswith(("dear ", "kind regards"))]
    if len(paragraphs) < 2:
        raise ValueError("Groq did not return a full letter.")
    return paragraphs


def cover_letter_stem(_full_name, role_family, company=None, job_title=None):
    """'Data Engineer_Bluecrux_CL': the cleaned job title, the company and the cover letter key."""
    return document_stem(job_title or role_family, company or "", "cover_letter")


def create_cover_letter_docx(role_family, company=None, job_title=None, data=None, output_path=None, location="",
                             paragraphs=None):
    """Build a cover letter, optionally addressed to one job, in the layout of the owner's sample letter.

    `paragraphs` is the body (see written_paragraphs); without it the role family's written letter is used.
    """
    letters = role_letters()
    if role_family not in letters and not company:
        raise ValueError(f"Unknown role family: {role_family}")

    data = data if data is not None else load_resume_file() or {}
    info = data.get("personal_info", {})
    paragraphs = paragraphs or fallback_paragraphs(role_family, data)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)

    def add(text="", size=11, bold=False, color=None, after=8, spacing=None):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.1
        if text:
            run = paragraph.add_run(text)
            run.font.name = FONT
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
            run.font.size = Pt(size)
            run.font.bold = bold
            if color:
                run.font.color.rgb = color
            if spacing:  # letter spacing, in twentieths of a point
                character_spacing = OxmlElement("w:spacing")
                character_spacing.set(qn("w:val"), str(spacing))
                run._element.get_or_add_rPr().append(character_spacing)
        return paragraph

    def add_rule(paragraph):
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "12"), ("space", "6"), ("color", "1F3A5F")):
            bottom.set(qn(f"w:{key}"), value)
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)

    add((info.get("full_name") or "").upper(), size=22, bold=True, color=NAVY, after=2, spacing=30)
    contact = [info.get(key) for key in ("location", "phone", "email") if info.get(key)]
    add_rule(add("   •   ".join(contact), size=9.5, color=MUTED, after=16))

    today = date.today()
    add(f"{today.day} {today:%B %Y}", color=MUTED, after=12)
    add("Hiring Team", after=4)
    addressee = ", ".join(part for part in (company, place_of(location, job_title)) if part and part != "the hiring team")
    add(addressee or "[Company Name]", after=14)
    add(f"Application: {clean_job_title(job_title, company) if job_title else role_family}", bold=True, color=NAVY,
        after=12)
    add("Dear Hiring Team,", after=8)
    for text in paragraphs:
        add(text, after=8)
    add("Kind regards,", after=4).paragraph_format.space_before = Pt(4)
    add(info.get("full_name") or "", bold=True, color=NAVY, after=0)

    if output_path:
        doc.save(output_path)
        return str(output_path)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = cover_letter_stem(info.get("full_name"), role_family, company, job_title)
    path = f"output/{stem}.docx"
    doc.save(long_path(PROJECT_ROOT / path))
    log(f"Cover letter created: {path}")
    return path


if __name__ == "__main__":
    for role in role_letters():
        create_cover_letter_docx(role)
