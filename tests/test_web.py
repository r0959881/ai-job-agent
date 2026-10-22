"""The hosted web app: sign-in, sessions and CSRF, per-user data, API keys, the run endpoints and documents."""
import base64
import io
from urllib.parse import parse_qs, urlparse

import pytest
from conftest import sign_in
from helpers import FakeGroq, fake_scrape, fake_site_search, resume_text, set_secret, user_id_for

from api import store


def new_client():
    from fastapi.testclient import TestClient

    from api.app import create_app

    return TestClient(create_app(), base_url="http://127.0.0.1")


# ---------------------------------------------------------------------------------- sessions

def test_health_and_session_before_sign_in(anon_client):
    assert anon_client.get("/api/health").json()["database"] is True
    session = anon_client.get("/api/session").json()
    assert session["user"] is None and session["csrf"]
    assert {"id": "dev", "label": "Developer sign-in"} in session["providers"]


def test_the_api_needs_a_signed_in_user(anon_client):
    response = anon_client.get("/api/resume")
    assert response.status_code == 401 and response.json()["code"] == "signed_out"
    # The failed request must not cost the CSRF token the sign-in needs next.
    sign_in(anon_client)
    assert anon_client.get("/api/resume").status_code == 200


def test_writes_need_the_csrf_token(client):
    token = client.headers.pop("X-CSRF-Token")
    assert client.put("/api/settings/secrets", json={"values": {}}).status_code == 403
    assert client.put("/api/settings/secrets", json={"values": {}}, headers={"X-CSRF-Token": "wrong"}).status_code == 403
    assert client.put("/api/settings/secrets", json={"values": {}}, headers={"X-CSRF-Token": token}).status_code == 200


def test_unknown_hosts_are_refused():
    from fastapi.testclient import TestClient

    from api.app import create_app

    with TestClient(create_app(), base_url="http://attacker.example") as other:
        assert other.get("/api/health").status_code == 400


def test_unknown_api_paths_return_json_404(client):
    response = client.get("/api/nope")
    assert response.status_code == 404 and "No API endpoint" in response.json()["detail"]


def test_logout_ends_the_session(client):
    assert client.post("/api/auth/logout").json() == {"signed_out": True}
    assert client.get("/api/resume").status_code == 401


def test_a_new_user_starts_with_their_name_and_an_empty_profile(anon_client):
    user = sign_in(anon_client, "grace@example.com", "Grace Hopper")
    body = anon_client.get("/api/resume").json()
    assert body["data"]["personal_info"]["full_name"] == "Grace Hopper"
    assert body["data"]["personal_info"]["email"] == "grace@example.com"
    assert body["empty"] is True  # onboarding still runs
    assert user["providers"] == ["dev"]


def test_signup_can_be_limited_to_listed_emails(anon_client, monkeypatch):
    monkeypatch.setenv("ALLOWED_EMAILS", "boss@example.com, @team.example")
    csrf = anon_client.get("/api/session").json()["csrf"]
    refused = anon_client.post("/api/auth/dev/login", json={"email": "x@other.example"}, headers={"X-CSRF-Token": csrf})
    assert refused.status_code == 403 and "invite-only" in refused.json()["detail"]
    sign_in(anon_client, "anyone@team.example")


# ------------------------------------------------------------------------------- isolation

def test_users_only_see_their_own_data(client):
    settings = client.get("/api/settings/search").json()["settings"]
    assert client.put("/api/settings/search", json={"settings": {**settings, "languages": ["nl"]}}).status_code == 200
    assert client.put("/api/settings/secrets", json={"values": {"GROQ_API_KEY": "gsk_mine_1234567"}}).status_code == 200

    other = new_client()
    sign_in(other, "other@example.com", "Other Person")
    assert other.get("/api/settings/search").json()["settings"]["languages"] == ["en"]
    assert other.get("/api/integrations/status").json()["groq"]["status"] == "missing"
    assert other.get("/api/resume").json()["data"]["personal_info"]["full_name"] == "Other Person"
    assert client.get("/api/resume").json()["data"]["personal_info"]["full_name"] == "Ammar Haider"


def test_api_keys_are_masked_and_encrypted_at_rest(client):
    response = client.put("/api/settings/secrets", json={"values": {"GROQ_API_KEY": "gsk_abcdefghijkl1234"}})
    groq = next(item for item in response.json() if item["key"] == "GROQ_API_KEY")
    assert groq == {"key": "GROQ_API_KEY", "label": "Groq API key", "group": "groq", "set": True, "masked": "gsk_…1234"}
    import sqlalchemy as sa

    with store.engine().connect() as conn:
        raw = conn.execute(sa.select(store.user_secrets.c.value)).scalar_one()
    assert "gsk_" not in raw
    assert store.get_secret(user_id_for(), "GROQ_API_KEY") == "gsk_abcdefghijkl1234"
    client.put("/api/settings/secrets", json={"values": {"GROQ_API_KEY": ""}})
    assert store.get_secret(user_id_for(), "GROQ_API_KEY") == ""
    assert client.put("/api/settings/secrets", json={"values": {"OTHER": "x"}}).status_code == 400


def test_integration_tests_explain_missing_keys(client):
    assert "Groq API key" in client.post("/api/integrations/groq/test").json()["message"]
    assert "Adzuna" in client.post("/api/integrations/adzuna/test").json()["message"]
    assert client.post("/api/integrations/gmail/test").status_code == 404  # Gmail is not part of the web app


def test_resume_history_lives_in_the_database(client):
    body = client.get("/api/resume").json()
    data = body["data"]
    data["languages"].append("French: A1")
    saved = client.put("/api/resume", json={"data": data, "etag": body["etag"]}).json()
    assert saved["backup_id"] and "French: A1" in resume_text()
    backups = client.get("/api/resume/backups").json()
    assert backups[0]["id"] == saved["backup_id"] and backups[0]["label"] == "Before saving in the editor"


def test_account_export_and_delete(client):
    set_secret("GROQ_API_KEY", "gsk_secret_value_1")
    export = client.get("/api/account/export")
    assert export.headers["content-disposition"].startswith("attachment")
    body = export.json()
    assert body["user"]["email"] == "tester@example.com" and "Ammar Haider" in body["resume_yaml"]
    assert body["api_keys_saved"] == ["GROQ_API_KEY"] and "gsk_secret_value_1" not in export.text

    assert client.post("/api/account/delete", json={"confirm": "delete"}).status_code == 400
    assert client.post("/api/account/delete", json={"confirm": "DELETE"}).json() == {"deleted": True}
    assert client.get("/api/resume").status_code == 401
    import sqlalchemy as sa

    with store.engine().connect() as conn:
        for table in (store.users, store.resumes, store.user_secrets, store.resume_versions):
            assert conn.execute(sa.select(sa.func.count()).select_from(table)).scalar_one() == 0


# ----------------------------------------------------------------------------------- OAuth

def test_google_sign_in_round_trip(anon_client, monkeypatch):
    import api.auth as auth

    monkeypatch.setenv("GOOGLE_CLIENT_ID", "gid")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "gsecret")
    monkeypatch.setenv("PUBLIC_URL", "https://jobs.example.com")
    redirect = anon_client.get("/api/auth/google/login?next=/jobs", follow_redirects=False)
    target = urlparse(redirect.headers["location"])
    query = parse_qs(target.query)
    assert target.netloc == "accounts.google.com"
    assert query["redirect_uri"] == ["https://jobs.example.com/api/auth/google/callback"]
    assert query["scope"] == ["openid email profile"] and query["code_challenge_method"] == ["S256"]

    seen = {}

    def fake_exchange(provider, code, redirect_uri, verifier=None):
        seen.update(provider=provider, code=code, verifier=verifier)
        return "google-sub-1", "Grace@Example.com", "Grace Hopper", "https://photo"

    monkeypatch.setattr(auth, "exchange_code", fake_exchange)
    forged = anon_client.get("/api/auth/google/callback?code=c&state=forged", follow_redirects=False)
    assert forged.headers["location"].startswith("/login?error=")

    redirect = anon_client.get("/api/auth/google/login?next=/jobs", follow_redirects=False)
    state = parse_qs(urlparse(redirect.headers["location"]).query)["state"][0]
    done = anon_client.get(f"/api/auth/google/callback?code=the-code&state={state}", follow_redirects=False)
    assert done.headers["location"] == "/jobs"
    assert seen["code"] == "the-code" and seen["verifier"]
    me = anon_client.get("/api/session").json()["user"]
    assert me["email"] == "grace@example.com" and me["providers"] == ["google"]


def test_linkedin_with_the_same_email_signs_into_the_same_account(anon_client, monkeypatch):
    user_a, _ = store.sign_in("google", "g-1", "same@example.com", "Same Person")
    user_b, created = store.sign_in("linkedin", "li-1", "same@example.com", "Same Person")
    assert user_a["id"] == user_b["id"] and not created
    assert sorted(store.providers_for(user_a["id"])) == ["google", "linkedin"]


def test_sign_in_rejects_outside_next_targets(anon_client, monkeypatch):
    monkeypatch.setenv("LINKEDIN_CLIENT_ID", "lid")
    monkeypatch.setenv("LINKEDIN_CLIENT_SECRET", "lsecret")
    anon_client.get("/api/auth/linkedin/login?next=https://evil.example", follow_redirects=False)
    session_next = "/"  # stored as "/" because the target is not a local path
    from api.auth import safe_next

    assert safe_next("https://evil.example") == session_next and safe_next("//evil.example") == "/"
    assert safe_next("/jobs?x=1") == "/jobs?x=1"


def test_unconfigured_provider_explains(anon_client):
    redirect = anon_client.get("/api/auth/google/login", follow_redirects=False)
    assert "not+set+up" in redirect.headers["location"]


# ------------------------------------------------------------------------------ search run

@pytest.fixture
def offline(monkeypatch):
    import discovery
    import tailor
    from job_sources import scrapers, sites

    monkeypatch.setattr(scrapers, "_scrape", fake_scrape)
    monkeypatch.setattr(sites, "search", fake_site_search)
    groq = FakeGroq()
    monkeypatch.setattr(tailor, "get_client", lambda: groq)
    monkeypatch.setattr(discovery, "adzuna_credentials", lambda: None)
    return groq


SCOPE = {"regions": ["BE"], "roles": ["Data Analyst"], "sources": ["linkedin", "vdab", "adzuna"], "languages": ["en"]}


def test_plan_and_run_return_jobs_without_storing_them(client, offline):
    plan = client.post("/api/search/plan", json=SCOPE).json()
    assert [(s["source"], s["label"]) for s in plan["searches"]][:2] == [
        ("linkedin", "Belgium · Data Analyst"), ("vdab", "Belgium · data analyst")]
    assert plan["skipped"] == [{"source": "adzuna", "label": "Adzuna", "kind": "api",
                                "reason": "add your Adzuna keys in Settings → Integrations"}]
    jobs = []
    for search in plan["searches"]:
        result = client.post("/api/search/run", json=SCOPE | {"index": search["index"], "source": search["source"],
                                                               "label": search["label"]}).json()
        assert result["ok"] and result["counts"]["kept"] == len(result["jobs"])
        jobs += result["jobs"]
    assert {job["role_family"] for job in jobs} == {"Data Analyst"}
    assert all(job["id"] and job["discovered_at"] for job in jobs)
    import sqlite3

    import config

    assert not config.DB_PATH.exists() or not sqlite3.connect(config.DB_PATH).execute(
        "SELECT name FROM sqlite_master WHERE name = 'jobs'").fetchone()


def test_a_changed_plan_is_refused(client, offline):
    assert client.post("/api/search/run", json=SCOPE | {"index": 99}).status_code == 409
    assert client.post("/api/search/run", json=SCOPE | {"index": 0, "label": "Somewhere else"}).status_code == 409


def test_a_failing_source_is_reported_not_raised(client, offline, monkeypatch):
    from job_sources import sites

    def broken(*_args, **_kwargs):
        raise ConnectionError("blocked")

    monkeypatch.setattr(sites, "search", broken)
    plan = client.post("/api/search/plan", json=SCOPE).json()
    vdab = next(s for s in plan["searches"] if s["source"] == "vdab")
    result = client.post("/api/search/run", json=SCOPE | {"index": vdab["index"]}).json()
    assert result == {"ok": False, "error": "ConnectionError: blocked", "jobs": [], "excluded": [],
                      "counts": {"discovered": 0, "kept": 0, "filtered": 0, "language_filtered": 0}}


def test_jobs_in_other_languages_come_back_as_excluded(client, offline, monkeypatch):
    from job_sources import sites

    def dutch(source, query, ctx, region, _cache=None):
        return [{"title": f"Junior {query}", "company_name": "NL BV", "location": "Gent, Belgium",
                 "url": "https://x", "description": "Voor onze klant zoeken we een junior data analist. Je werkt met "
                                                    "de teams en je bent sterk in SQL en Power BI rapporten."}]

    monkeypatch.setattr(sites, "search", dutch)
    plan = client.post("/api/search/plan", json=SCOPE).json()
    vdab = next(s for s in plan["searches"] if s["source"] == "vdab")
    result = client.post("/api/search/run", json=SCOPE | {"index": vdab["index"]}).json()
    assert result["jobs"] == [] and result["counts"]["language_filtered"] == 1
    assert result["excluded"][0]["exclusion_reason"].startswith("Posted in Dutch")


# ------------------------------------------------------------------------------- documents

JOB = {"id": "j1", "title": "Junior Data Analyst", "company": "Acme NV", "location": "Gent, Belgium",
       "source": "LinkedIn", "description": "Junior analyst with SQL, Python and Power BI."}


def test_tailor_needs_the_users_groq_key(client):
    response = client.post("/api/tailor", json={"job": JOB})
    assert response.status_code == 400 and "Groq API key" in response.json()["detail"]


def test_tailor_returns_the_cv_file(client, offline):
    preview = client.post("/api/tailor/preview", json={"job": JOB}).json()
    assert preview["summary"] == "Tailored summary." and preview["role_family"] == "Data Analyst"
    result = client.post("/api/tailor", json={"job": JOB, "summary": "Edited.", "bullets": ["Mine"]}).json()
    assert result["file"]["name"] == "Junior Data Analyst_Acme_CV.pdf"
    assert result["file"]["mime"] == "application/pdf"
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(base64.b64decode(result["file"]["data"])))
    text = reader.pages[0].extract_text()
    assert "Edited." in text and "Mine" in text and len(reader.pages) <= 2
    assert offline.completions.calls == 1  # the second call used the edited content


def test_cover_letter_report_preview_and_insights(client, offline):
    letter = client.post("/api/cover-letter", json={"job": JOB}).json()
    assert letter["file"]["name"] == "Junior Data Analyst_Acme_CL.docx"
    report = client.post("/api/report", json={"jobs": [JOB | {"cv_path": "../CVs/cv.docx"}], "excluded": []}).json()
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(base64.b64decode(report["file"]["data"])))["Jobs"]
    assert sheet.cell(2, 15).hyperlink.target == "../CVs/cv.docx"

    letter_file = base64.b64decode(letter["file"]["data"])
    html = client.post("/api/documents/preview", files={"file": ("letter.docx", letter_file)}).json()["html"]
    assert "AMMAR HAIDER" in html and "<script" not in html
    assert "Application: Junior Data Analyst" in html and "Dear Hiring Team," in html and "Kind regards," in html
    assert client.post("/api/documents/preview", files={"file": ("x.pdf", b"%PDF")}).status_code == 415

    insights = client.post("/api/jobs/insights", json=JOB).json()
    assert "SQL" in insights["matched_keywords"] and insights["tips"]


def test_an_empty_profile_cannot_be_tailored(anon_client, offline):
    sign_in(anon_client, "new@example.com", "New Person")
    response = anon_client.post("/api/tailor", json={"job": JOB})
    assert response.status_code == 400 and "profile is empty" in response.json()["detail"]


# ------------------------------------------------------------------------------ local mode

@pytest.fixture
def local_client(monkeypatch, project_root):
    monkeypatch.setenv("AI_JOB_AGENT_LOCAL", "1")
    client = new_client()
    session = client.get("/api/session").json()
    client.headers["X-CSRF-Token"] = session["csrf"]
    client.session_info = session
    return client


def test_local_mode_signs_in_to_the_project_files(local_client, project_root, monkeypatch):
    session = local_client.session_info
    assert session["project_files"] is True and session["user"]["name"] == "Ammar Haider"
    assert local_client.get("/api/resume").json()["data"]["personal_info"]["full_name"] == "Ammar Haider"
    # Keys go to .env, not the database
    (project_root / ".env").write_text("# keep me\nGROQ_API_KEY=old\n", encoding="utf-8")
    monkeypatch.setenv("GROQ_API_KEY", "old")
    local_client.put("/api/settings/secrets", json={"values": {"GROQ_API_KEY": "gsk_local_key_12345"}})
    text = (project_root / ".env").read_text(encoding="utf-8")
    assert "# keep me" in text and "GROQ_API_KEY=gsk_local_key_12345" in text
    assert local_client.get("/api/integrations/status").json()["groq"]["status"] == "ok"
    # Settings go to agent_settings.yaml
    settings = local_client.get("/api/settings/search").json()["settings"]
    local_client.put("/api/settings/search", json={"settings": {**settings, "languages": ["nl"]}})
    assert "- nl" in (project_root / "agent_settings.yaml").read_text(encoding="utf-8")


def test_local_mode_shares_jobs_db_and_output(local_client, project_root):
    from database import init_db, transaction

    init_db()
    with transaction() as conn:
        conn.execute("INSERT INTO jobs (id, title, company, location, cv_path, application_status) "
                     "VALUES ('j1', 'Junior Data Engineer', 'Acme', 'Gent', 'output/Me_CV_Acme_Junior.docx', 'applied')")
    (project_root / "output" / "Me_CV_Acme_Junior.docx").write_bytes(b"cv")
    (project_root / "output" / "job_report_1.xlsx").write_bytes(b"xlsx")
    legacy = local_client.get("/api/local/legacy").json()
    assert legacy["jobs"][0]["application_status"] == "applied"
    by_path = {doc["path"]: doc for doc in legacy["documents"]}
    assert by_path["output/Me_CV_Acme_Junior.docx"]["job_id"] == "j1"
    assert by_path["output/job_report_1.xlsx"]["kind"] == "report"

    assert local_client.put("/api/local/files/new.docx", content=b"new").status_code == 200
    assert (project_root / "output" / "new.docx").read_bytes() == b"new"
    assert local_client.get("/api/local/files/new.docx").content == b"new"
    assert local_client.put("/api/local/files/..%5Cjobs.db", content=b"x").status_code in {400, 404}
    assert local_client.delete("/api/local/files/new.docx").status_code == 200
    assert not (project_root / "output" / "new.docx").exists()


def test_signing_out_of_the_local_account_sticks(local_client):
    local_client.post("/api/auth/logout")
    session = local_client.get("/api/session").json()
    assert session["user"] is None and {"id": "local", "label": "Local account (this PC's project files)"} in session["providers"]
    local_client.headers["X-CSRF-Token"] = session["csrf"]
    assert local_client.post("/api/auth/local/login").status_code == 200
    assert local_client.get("/api/session").json()["project_files"] is True


def test_project_files_are_not_available_to_web_users(client):
    assert client.get("/api/local/legacy").status_code == 404
    assert client.get("/api/session").json()["project_files"] is False


def test_the_owner_email_opens_the_project_files(monkeypatch, project_root):
    monkeypatch.setenv("AI_JOB_AGENT_LOCAL", "1")
    monkeypatch.setenv("AI_JOB_AGENT_LOCAL_EMAIL", "Owner@Example.com")
    automatic = new_client().get("/api/session").json()
    assert automatic["project_files"] and automatic["user"]["email"] == "owner@example.com"

    owner = new_client()
    user = sign_in(owner, "owner@example.com")
    assert user["id"] == automatic["user"]["id"]
    assert owner.get("/api/session").json()["project_files"] is True
    assert owner.get("/api/resume").json()["data"]["personal_info"]["full_name"] == "Ammar Haider"

    someone_else = new_client()
    sign_in(someone_else, "tester2@example.com", "Tester Two")
    assert someone_else.get("/api/session").json()["project_files"] is False
    assert someone_else.get("/api/resume").json()["data"]["personal_info"]["full_name"] == "Tester Two"


def test_the_cv_plan_is_tailored_and_fits_two_pages():
    import copy

    import yaml
    from conftest import CODE_ROOT
    from pypdf import PdfReader

    from generate_doc import FONT, build_cv_pdf, build_plan, split_tag

    data = yaml.safe_load((CODE_ROOT / "resume_data.yaml").read_text(encoding="utf-8"))
    assert FONT == "Carlito"  # the bundled standard font is found
    assert split_tag("[2] [2] Built **ETL**") == (1, "Built **ETL**")
    job = "We need a data engineer with SQL, Python and Power BI."
    plan = build_plan(data, "Data Engineer", "Summary.", ["[2] Second role bullet", "[1] First role bullet"],
                      tagline="Data engineer for Acme", skills=["SQL", "Not In Resume", "Python"],
                      job_title="Data Engineer", company="Acme", job_text=job)
    assert plan.tagline == "Data engineer for Acme"
    assert [bullets for _entry, bullets in plan.experience][:2] == [["First role bullet"], ["Second role bullet"]]
    assert "Not In Resume" not in plan.skills and {"SQL", "Python", "Power BI"} <= set(plan.bold_terms)
    preferred = data["role_profiles"]["Data Engineer"]["preferred_projects"]
    titles = [p["title"] for p, _h in plan.projects]
    assert titles == preferred  # only the projects that fit the role

    pdf = build_cv_pdf(data, "Acme", "Data Engineer", "Summary.", ["[1] First role bullet"], "Data Engineer",
                       tagline="Data engineer for Acme", job_text=job)
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) <= 2
    assert "Dutch: A2" in reader.pages[-1].extract_text().splitlines()[-1]  # languages close the CV

    heavy = copy.deepcopy(data)
    heavy["experience"] = heavy["experience"] * 3
    pdf = build_cv_pdf(heavy, "Acme", "Data Engineer", "A long summary sentence. " * 12, [], "Data Engineer",
                       job_text=job)
    assert len(PdfReader(io.BytesIO(pdf)).pages) <= 2  # trimmed as far as it goes; one page for real profiles


def test_tailored_text_with_numbers_not_in_the_resume_is_left_out():
    import tailor

    resume = {"experience": [{"bullets": ["Cut setup time by 80% for 50+ users."]}]}
    assert tailor.keep_verifiable(["[1] Cut setup by 80%", "[1] Covered 75 sources", "[2] No numbers"], resume) == [
        "[1] Cut setup by 80%", "[2] No numbers"]
    summary = "I bring 1.5 years of experience. I cut setup time by 80% for 50+ users."
    assert tailor.drop_unverifiable_sentences(summary, resume) == "I cut setup time by 80% for 50+ users."


def test_autofill_is_local_only_and_checks_the_link(local_client, client):
    assert client.post("/api/local/autofill", json={"url": "https://example.com"}).status_code == 404
    response = local_client.post("/api/local/autofill", json={"url": "javascript:alert(1)"})
    assert response.status_code == 400 and "not a web address" in response.json()["detail"]
    from api.routers.local import _output_file

    assert _output_file("output/../jobs.db") == "" and _output_file("../secret.txt") == ""
