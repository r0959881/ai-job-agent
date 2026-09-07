import os
import yaml
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_tailored_cv_docx(company_name, job_title, tailored_summary, tailored_bullets, role_family=None):
    """Generates a professional .docx CV formatted for ATS parsers."""
    
    # Load master profile info
    if not os.path.exists("resume_data.yaml"):
        print("❌ Error: resume_data.yaml not found!")
        return None

    with open("resume_data.yaml", "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)

    role_profile = data.get("role_profiles", {}).get(role_family or "", {})
        
    info = data.get("personal_info", {})
    doc = Document()
    
    navy = RGBColor(31, 62, 99)
    body_font = "Times New Roman"

    # Match the reference PDF's compact, single-column layout.
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = body_font
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    normal_style.font.size = Pt(10)

    def set_run_font(run, size=10, bold=False, italic=False, color=None):
        run.font.name = body_font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color

    def set_rule(paragraph):
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), "B7B7B7")
        borders.append(bottom)
        p_pr.append(borders)
        
    # --- HEADER ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(1)
    run_name = title_p.add_run(info.get("full_name", "Ammar Haider").upper())
    set_run_font(run_name, size=19, bold=True, color=navy)
    
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    contact_p.paragraph_format.space_after = Pt(4)
    
    # Build contact string safely without KeyError
    contact_parts = [
        info.get("location"),
        info.get("phone"),
        info.get("email"),
        info.get("work_status") or info.get("work_authorization")
    ]
    contact_str = " | ".join([part for part in contact_parts if part])
    
    set_run_font(contact_p.add_run(info.get("headline", "")), size=10.5)
    contact_p.add_run("\n")
    set_run_font(contact_p.add_run(contact_str), size=9.5, color=RGBColor(80, 80, 80))
    
    # Helper to add section headings
    def add_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text.upper())
        set_run_font(run, size=10.5, bold=True, color=navy)
        set_rule(h)
        
    # --- PROFESSIONAL SUMMARY ---
    add_heading("Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    set_run_font(p_sum.add_run(tailored_summary), size=10)

    if role_profile.get("keywords"):
        add_heading("Key Skills Relevant to This Role")
        p_role = doc.add_paragraph(style="List Bullet")
        p_role.paragraph_format.space_after = Pt(3)
        set_run_font(p_role.add_run(", ".join(role_profile["keywords"])), size=10)

    # --- CORE SKILLS ---
    add_heading("Technical Skills & Core Competencies")
    skills = data.get("core_skills", {})
    preferred_categories = role_profile.get("preferred_skill_categories", [])
    if preferred_categories:
        skills = {category: skills[category] for category in preferred_categories if category in skills}
    if isinstance(skills, dict):
        for category, skill_list in skills.items():
            p_sk = doc.add_paragraph()
            p_sk.paragraph_format.space_after = Pt(2)
            set_run_font(p_sk.add_run(f"• {category.replace('_', ' ').title()}: "), size=9.5, bold=True)
            set_run_font(p_sk.add_run(", ".join(skill_list) if isinstance(skill_list, list) else str(skill_list)), size=9.5)
    elif isinstance(skills, list):
        p_sk = doc.add_paragraph()
        p_sk.add_run("• Skills: " + ", ".join(skills)).font.size = Pt(9.5)

    # --- PROFESSIONAL EXPERIENCE ---
    add_heading("Professional Experience")
    
    if tailored_bullets:
        p_tb_head = doc.add_paragraph()
        r_tb = p_tb_head.add_run("Key Role-Matched Highlights:")
        r_tb.font.bold = True
        r_tb.font.size = Pt(9.5)
        for bullet in tailored_bullets:
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_after = Pt(2)
            set_run_font(p_b.add_run(bullet), size=9.5)

    for exp in data.get("experience", []):
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_before = Pt(4)
        p_exp.paragraph_format.space_after = Pt(2)
        p_exp.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
        set_run_font(p_exp.add_run(exp.get("role", "")), size=10, bold=True)
        p_exp.add_run("\t")
        set_run_font(p_exp.add_run(exp.get("dates", "")), size=9.5, bold=True)

        p_company = doc.add_paragraph()
        p_company.paragraph_format.space_after = Pt(1)
        set_run_font(p_company.add_run(f"{exp.get('company', '')} — {exp.get('location', '')}"), size=9.5, italic=True, color=RGBColor(80, 80, 80))
        
        for bullet in exp.get("bullets", []):
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_after = Pt(2)
            set_run_font(p_b.add_run(bullet), size=9.5)

    # --- PROJECTS ---
    projects = data.get("projects", [])
    preferred_projects = role_profile.get("preferred_projects", [])
    if preferred_projects:
        projects = [project for project in projects if project.get("title") in preferred_projects]
    if projects:
        add_heading("Relevant Projects")
        for project in projects:
            p_project = doc.add_paragraph()
            p_project.paragraph_format.space_before = Pt(4)
            p_project.paragraph_format.space_after = Pt(2)

            set_run_font(p_project.add_run(project.get("title", "")), size=10, bold=True)

            tech_stack = project.get("tech_stack", "")
            if tech_stack:
                set_run_font(p_project.add_run(f"\n{tech_stack}"), size=9.5, italic=True, color=RGBColor(80, 80, 80))

            for highlight in project.get("highlights", []):
                p_highlight = doc.add_paragraph(style='List Bullet')
                p_highlight.paragraph_format.space_after = Pt(2)
                set_run_font(p_highlight.add_run(highlight), size=9.5)

    # --- EDUCATION ---
    add_heading("Education")
    edu = data.get("education", {})
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_after = Pt(2)
    
    if isinstance(edu, dict):
        set_run_font(p_edu.add_run(f"• {edu.get('degree', '')} — {edu.get('institution', '')} ({edu.get('dates', edu.get('year', ''))})"), size=9.5)
        coursework = edu.get("coursework", [])
        if coursework:
            p_coursework = doc.add_paragraph()
            p_coursework.paragraph_format.space_after = Pt(2)
            set_run_font(p_coursework.add_run("Key coursework: " + ", ".join(coursework)), size=9.5)
    elif isinstance(edu, list):
        for e in edu:
            p_e = doc.add_paragraph()
            p_e.paragraph_format.space_after = Pt(2)
            set_run_font(p_e.add_run(f"• {e.get('degree', '')} — {e.get('institution', '')} ({e.get('year', '')})"), size=9.5)

    # --- CERTIFICATIONS ---
    add_heading("Certifications")
    for cert in data.get("certifications", []):
        p_cert = doc.add_paragraph()
        p_cert.paragraph_format.space_after = Pt(2)
        if isinstance(cert, dict):
            set_run_font(p_cert.add_run(f"• {cert.get('name', '')} ({cert.get('issuer', '')}, {cert.get('year', '')})"), size=9.5)
        else:
            set_run_font(p_cert.add_run(f"• {cert}"), size=9.5)

    # Save Document
    os.makedirs("output", exist_ok=True)
    clean_company = "".join(c for c in str(company_name) if c.isalnum() or c in (' ', '_')).strip()
    clean_title = "".join(c for c in str(job_title) if c.isalnum() or c in (' ', '_')).strip()
    # Keep generated paths below Windows' maximum path length.
    filename_prefix = "Ammar_Haider_CV_"
    max_stem_length = 180 - len(filename_prefix)
    filename_stem = f"{filename_prefix}{clean_company}_{clean_title}"[:max_stem_length]
    filename_stem = filename_stem.rstrip(" ._")
    
    file_path = f"output/{filename_stem}.docx"
    doc.save(file_path)
    print(f"📄 Word Document created successfully: {file_path}")
    return file_path

if __name__ == "__main__":
    sample_summary = "Data Engineer with proven expertise building Medallion architectures and Power BI solutions across Azure Databricks and Microsoft Fabric."
    sample_bullets = [
        "Designed and implemented a production-grade ETL pipeline that ingested >50K IoT sensor records daily into Delta Lake.",
        "Built a CI/CD framework with GitHub Actions, Docker, and Kubernetes reducing manual setup time by 80%."
    ]
    create_tailored_cv_docx("Algo1", "Software_Engineer_Platform", sample_summary, sample_bullets)